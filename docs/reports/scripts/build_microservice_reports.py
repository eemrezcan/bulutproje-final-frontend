from __future__ import annotations

from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from PIL import Image as PILImage, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "assets"
OUT_DIR = ROOT / "out"

PUBLIC_URL = "https://d1k4ey8nhx4dtb.cloudfront.net"
FRONTEND_REPO = "https://github.com/eemrezcan/bulutproje-final-frontend"

STUDENT = "Ahmet Emre Özcan"
STUDENT_NO = "21290560"
COURSE = "3522 Bulut Bilişim"
INSTRUCTOR = "Enver Bağcı"
PROJECT = "Akıllı Şehir Yönetim Platformu"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(91, 103, 118)
LIGHT_FILL = "F2F4F7"
RULE = "DADCE0"


@dataclass
class ReportSpec:
    slug: str
    title: str
    scope: str
    module_name: str
    repo_label: str
    repo_url: str
    demo_video_url: str
    screenshot: str
    abstract: str
    keywords: list[str]
    pdf_project: str
    pdf_expectations: list[str]
    service_boundary_do: list[str]
    service_boundary_dont: list[str]
    technologies: list[tuple[str, str]]
    local_flow: str
    aws_flow: str
    data_tables: list[tuple[str, list[tuple[str, str]]]]
    endpoints: list[str]
    deployed_resources: list[tuple[str, str]]
    validation: list[tuple[str, str]]
    challenges: list[tuple[str, str]]
    conclusion: str
    sources: list[str]


def set_run_font(run, name="Noto Sans", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[min(idx, len(widths_dxa) - 1)]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def cover_summary_rows(spec: ReportSpec) -> list[tuple[str, str]]:
    aws_services = ", ".join(name for name, _ in spec.deployed_resources[:4])
    return [
        ("Kapsam", f"{spec.pdf_project}; {spec.scope}"),
        ("AWS odağı", aws_services),
        (
            "Mikroservis sınırı",
            "Ayrı API, ayrı veri modeli, ayrı deployment ve dar IAM yetki sınırı",
        ),
    ]


def add_paragraph(doc: Document, text: str = "", style: str | None = None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix) :])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r)


def local_aws_rows(spec: ReportSpec) -> list[list[str]]:
    if spec.module_name == "IoT Service":
        return [
            ["Mosquitto MQTT Broker", "AWS IoT Core"],
            ["IoT MQTT consumer", "IoT Rule + Lambda ingestion"],
            ["DynamoDB Local", "Amazon DynamoDB: iot_sensors, iot_sensor_readings"],
            ["FastAPI IoT container", "ECS Fargate IoT service"],
            ["Local simulator", "IoT Core topic publish veya cihaz/simulator mesajı"],
        ]
    if spec.module_name == "Video Analysis Service":
        return [
            ["Local sentetik video sahnesi", "S3 video object ve ECS MJPEG canlı demo endpointi"],
            ["OpenCV/sentetik analiz", "Rekognition Video stored analysis + Lambda result processing"],
            ["DynamoDB Local", "Amazon DynamoDB: video_cameras, video_analysis_results"],
            ["FastAPI Video container", "ECS Fargate Video service"],
            ["Local stream endpoint", "Kinesis Video Streams seçeneği veya CloudFront üzerinden canlı demo yolu"],
        ]
    return [
        ["Local joblib/scikit-learn model", "Lambda inference veya ölçekli alternatif olarak SageMaker endpoint"],
        ["Local REST veri okuma", "CloudFront path routing üzerinden IoT ve Video API okuma"],
        ["DynamoDB Local", "Amazon DynamoDB: ml_predictions"],
        ["FastAPI ML container", "ECS Fargate ML service"],
        ["Manual prediction üretimi", "Lambda invoke veya ML API prediction generate akışı"],
    ]


def cost_rows() -> list[list[str]]:
    return [
        ["Geçici kapatma", "ECS service desired-count=0 yapılarak Fargate task maliyeti durdurulabilir."],
        ["Veri koruma", "DynamoDB, S3, ECR ve IAM kaynakları silinmeden demo verileri korunur."],
        ["ALB notu", "ALB'nin stop modu yoktur; tamamen maliyeti kesmek için silmek gerekir, demo öncesi önerilmez."],
        ["Yaklaşık maliyet", "3 Fargate task + ALB + public IPv4 düşük trafikte yaklaşık 0.13 USD/saat düzeyindedir."],
    ]


def scope_mapping_rows(spec: ReportSpec) -> list[list[str]]:
    if spec.module_name == "IoT Service":
        return [
            ["Akıllı şehir sensörleri", "Meydan, Otogar, Kampus, Hastane ve Sanayi bölgeleri için sentetik sensör noktaları kullanıldı."],
            ["IoT veri toplama", "Sıcaklık, nem, hava kalitesi ve trafik yoğunluğu telemetry verisi üretildi."],
            ["IoT protokolü", "Telemetry verisi REST ile değil MQTT topic yapısı üzerinden sisteme alındı."],
            ["Bulut karşılığı", "AWS IoT Core, IoT Rule, Lambda ve DynamoDB hattı ile gerçek AWS ingestion akışı kuruldu."],
            ["Görselleştirme", "React panelde son ölçümler, bölge kartları ve zaman serisi grafiği gösterildi."],
        ]
    if spec.module_name == "Video Analysis Service":
        return [
            ["Akıllı şehir kamera noktaları", "Beş şehir bölgesi için ayrı kamera metadata'sı ve canlı görüntü alanı tanımlandı."],
            ["Video akışı", "ECS üzerindeki MJPEG endpoint ile canlı kamera hissi veren şehir sahnesi yayınlandı."],
            ["Video analizi", "Kişi sayısı, araç sayısı, hareket seviyesi, kalabalık seviyesi ve label eventleri üretildi."],
            ["Bulut karşılığı", "S3 üzerindeki video dosyası Rekognition Video stored analysis ile işlendi; SNS + Lambda sonucu DynamoDB'ye yazdı."],
            ["Veri saklama", "Video dosyası DynamoDB'de değil S3'te; analiz eventleri DynamoDB'de tutuldu."],
        ]
    return [
        ["Akıllı veri analitiği", "IoT ve Video servislerinden gelen özet veriler bölge bazında birleştirildi."],
        ["Makine öğrenmesi", "scikit-learn, pandas ve joblib ile sıcaklık/kalabalık tahmini üretildi."],
        ["Karar destek", "Rule-based risk motoru ile risk skoru, risk seviyesi ve öneri metni oluşturuldu."],
        ["Bulut karşılığı", "ML API ECS Fargate üzerinde çalıştı; düşük maliyetli inference için AWS Lambda kullanıldı."],
        ["Veri saklama", "Tahminler ve risk kayıtları yalnızca ml_predictions DynamoDB tablosunda tutuldu."],
    ]


def flow_rows(flow: str) -> list[list[str]]:
    rows: list[list[str]] = []
    cleaned = [line.strip() for line in flow.splitlines() if line.strip()]
    for idx, line in enumerate(cleaned, start=1):
        label = "Başlangıç" if idx == 1 else f"Adım {idx - 1}"
        rows.append([label, line.removeprefix("->").strip()])
    return rows


def source_summary(spec: ReportSpec) -> str:
    module_sources = {
        "IoT Service": "AWS IoT Core ve Mosquitto",
        "Video Analysis Service": "Rekognition Video, SNS, Kinesis Video Streams ve OpenCV",
        "ML Analytics Service": "scikit-learn, Pandas, Lambda ve SageMaker",
    }
    return (
        "Kaynaklar: 3522 Bulut Bilişim final proje yönergesi; AWS resmi "
        f"dokümantasyonu; FastAPI, React/Vite/TanStack Query/Recharts ve {module_sources[spec.module_name]} "
        "dokümantasyonları; proje repository dokümanları."
    )


def _text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont) -> tuple[int, int]:
    box = draw.multiline_textbbox((0, 0), text, font=font, spacing=8, align="center")
    return box[2] - box[0], box[3] - box[1]


def _draw_box(draw: ImageDraw.ImageDraw, xy, text: str, font, fill="#F0EEFF", outline="#9B7FE8"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=8, fill=fill, outline=outline, width=3)
    tw, th = _text_size(draw, text, font)
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2), text, font=font, fill="#2B2B35", spacing=8, align="center")


def _draw_db(draw: ImageDraw.ImageDraw, xy, text: str, font, fill="#F0EEFF", outline="#9B7FE8"):
    x1, y1, x2, y2 = xy
    h = y2 - y1
    ellipse_h = min(54, h // 4)
    draw.rectangle((x1, y1 + ellipse_h / 2, x2, y2 - ellipse_h / 2), fill=fill, outline=outline, width=3)
    draw.ellipse((x1, y1, x2, y1 + ellipse_h), fill=fill, outline=outline, width=3)
    draw.arc((x1, y2 - ellipse_h, x2, y2), 0, 180, fill=outline, width=3)
    draw.line((x1, y1 + ellipse_h / 2, x1, y2 - ellipse_h / 2), fill=outline, width=3)
    draw.line((x2, y1 + ellipse_h / 2, x2, y2 - ellipse_h / 2), fill=outline, width=3)
    tw, th = _text_size(draw, text, font)
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2 + 6), text, font=font, fill="#2B2B35", spacing=8, align="center")


def _arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]):
    draw.line((start, end), fill="#55545C", width=4)
    import math

    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    spread = 0.45
    p1 = (end[0] - length * math.cos(angle - spread), end[1] - length * math.sin(angle - spread))
    p2 = (end[0] - length * math.cos(angle + spread), end[1] - length * math.sin(angle + spread))
    draw.polygon([end, p1, p2], fill="#55545C")


def architecture_nodes(spec: ReportSpec):
    if spec.module_name == "IoT Service":
        return {
            "api": "IoT FastAPI\nBackend API",
            "db": "DynamoDB\nIoT Tabloları",
            "bottom": ["Sensor /\nSimulator", "AWS IoT Core", "IoT Rule\n+ Lambda"],
            "panel": "React IoT\nPaneli",
        }
    if spec.module_name == "Video Analysis Service":
        return {
            "api": "Video FastAPI\nBackend API",
            "db": "DynamoDB\nVideo Eventleri",
            "bottom": ["Amazon S3\nVideo", "Rekognition\nVideo", "SNS + Lambda"],
            "panel": "React Video\nPaneli",
        }
    return {
        "api": "ML FastAPI\nBackend API",
        "db": "DynamoDB\nML Tahminleri",
        "bottom": ["IoT API +\nVideo API", "Feature\nExtraction", "Lambda\nInference"],
        "panel": "React ML\nPaneli",
    }


def ensure_architecture_diagram(spec: ReportSpec) -> Path:
    out = ASSET_DIR / f"{spec.slug}-architecture.png"
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    w, h = 1780, 820
    img = PILImage.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf", 38)
    small = ImageFont.truetype("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf", 34)
    nodes = architecture_nodes(spec)

    boxes = {
        "user": (70, 80, 350, 190),
        "cf": (520, 80, 900, 190),
        "alb": (1010, 80, 1320, 190),
        "api": (1420, 80, 1710, 190),
        "admin": (70, 430, 350, 540),
        "panel": (520, 430, 900, 540),
        "b0": (90, 650, 360, 760),
        "b1": (520, 650, 830, 760),
        "b2": (980, 650, 1290, 760),
        "db": (1455, 570, 1690, 780),
    }
    _draw_box(draw, boxes["user"], "Kullanıcı", font)
    _draw_box(draw, boxes["cf"], "CloudFront\nWeb/API", font)
    _draw_box(draw, boxes["alb"], "ALB Path\nRouting", font)
    _draw_box(draw, boxes["api"], nodes["api"], small)
    _draw_box(draw, boxes["admin"], "Yönetici", font)
    _draw_box(draw, boxes["panel"], nodes["panel"], small)
    _draw_box(draw, boxes["b0"], nodes["bottom"][0], small)
    _draw_box(draw, boxes["b1"], nodes["bottom"][1], small)
    _draw_box(draw, boxes["b2"], nodes["bottom"][2], small)
    _draw_db(draw, boxes["db"], nodes["db"], small)

    _arrow(draw, (350, 135), (520, 135))
    _arrow(draw, (900, 135), (1010, 135))
    _arrow(draw, (1320, 135), (1420, 135))
    _arrow(draw, (350, 485), (520, 485))
    _arrow(draw, (900, 485), (1420, 170))
    _arrow(draw, (1420, 160), (900, 455))
    _arrow(draw, (360, 705), (520, 705))
    _arrow(draw, (830, 705), (980, 705))
    _arrow(draw, (1290, 705), (1455, 675))
    _arrow(draw, (1710, 135), (1568, 570))
    _arrow(draw, (710, 540), (710, 650))
    img.save(out)
    return out


def add_docx_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(header)
        set_run_font(r, bold=True, color=INK)
        set_cell_shading(cell, LIGHT_FILL)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=9.5)
    set_table_geometry(table, widths)
    return table


def setup_docx_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Noto Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        st = styles[style_name]
        st.font.name = "Noto Sans"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def set_header_footer(doc: Document, spec: ReportSpec):
    for section in doc.sections:
        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_p.text = ""
        r = header_p.add_run(f"{PROJECT} | {spec.module_name}")
        set_run_font(r, size=8.5, color=MUTED)
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = footer_p.add_run("Sayfa ")
        set_run_font(r, size=8.5, color=MUTED)
        add_page_number(footer_p)


def add_cover(doc: Document, spec: ReportSpec):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COURSE)
    set_run_font(r, size=13, color=MUTED, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(PROJECT)
    set_run_font(r, size=24, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(spec.title)
    set_run_font(r, size=18, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(spec.scope)
    set_run_font(r, size=12, color=MUTED, italic=True)

    doc.add_paragraph()
    rows = [
        ("Öğrenci", STUDENT),
        ("Öğrenci No", STUDENT_NO),
        ("Ders", COURSE),
        ("Öğretim Elemanı", INSTRUCTOR),
        ("Frontend Repo Linki", FRONTEND_REPO),
        (spec.repo_label, spec.repo_url),
        ("Demo Video Linki", spec.demo_video_url),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = ""
        cells[1].text = ""
        r = cells[0].paragraphs[0].add_run(label)
        set_run_font(r, bold=True, size=10)
        r = cells[1].paragraphs[0].add_run(value)
        set_run_font(r, size=10)
    set_table_geometry(table, [2500, 6860])

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Kapak proje özeti")
    set_run_font(r, size=11, color=BLUE, bold=True)
    summary_table = doc.add_table(rows=0, cols=2)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_table.style = "Table Grid"
    for label, value in cover_summary_rows(spec):
        cells = summary_table.add_row().cells
        cells[0].text = ""
        cells[1].text = ""
        r = cells[0].paragraphs[0].add_run(label)
        set_run_font(r, bold=True, size=9.5)
        r = cells[1].paragraphs[0].add_run(value)
        set_run_font(r, size=9.5)
    set_table_geometry(summary_table, [2500, 6860])
    doc.add_page_break()


def add_front_matter(doc: Document, spec: ReportSpec):
    doc.add_heading("Özet", level=1)
    add_paragraph(doc, spec.abstract)
    add_paragraph(doc, "Anahtar Kelimeler: " + ", ".join(spec.keywords))
    doc.add_heading("İçindekiler", level=1)
    toc = [
        "1. Giriş",
        "2. Kapsam",
        "3. Mikroservis sınırı ve sistem mimarisi",
        "4. Teknoloji kararları",
        "5. Veri modeli ve API tasarımı",
        "6. Local çalışma ve demo akışı",
        "7. AWS deployment mimarisi",
        "8. Güvenlik, IAM ve veri sahipliği",
        "9. Test, smoke doğrulama ve kanıtlar",
        "10. Ortak platform entegrasyonu",
        "11. Karşılaşılan sorunlar ve çözüm kararları",
        "12. Sonuç",
        "Kaynaklar",
    ]
    add_numbered(doc, toc)
    doc.add_page_break()


def add_module_sections(doc: Document, spec: ReportSpec):
    doc.add_heading("1. Giriş", level=1)
    add_paragraph(
        doc,
        f"Bu rapor, {PROJECT} kapsamında geliştirilen {spec.module_name} bileşenini "
        "ayrı bir mikroservis olarak ele alır. Platform tek bir yönetim paneli gibi "
        "kullanılsa da her mikroservis kendi veri modeline, kendi API yüzeyine, kendi "
        "container image'ına ve kendi AWS yetki sınırına sahiptir.",
    )
    add_paragraph(
        doc,
        "Temel tasarım hedefi, ders kapsamındaki bulut servislerini yalnızca isim olarak "
        "kullanmak değil, doğru problemi doğru AWS servisi ile eşleştirmektir. Bu nedenle "
        "raporda local demo akışı ile AWS deployment karşılığı birlikte anlatılmıştır.",
    )

    doc.add_heading("2. Kapsam", level=1)
    add_paragraph(
        doc,
        f"Bu bölüm, {spec.module_name} mikroservisinin proje içindeki sınırını ve "
        f"{COURSE} final yönergesindeki {spec.pdf_project} beklentilerini nasıl "
        "karşıladığını özetler. Kapsam, yalnızca uygulama ekranını değil; veri giriş "
        "kanalını, servis sorumluluğunu, AWS karşılığını ve raporda savunulacak "
        "deployment kararlarını birlikte ele alır.",
    )
    add_docx_table(
        doc,
        ["Kapsam maddesi", "Projede karşılığı"],
        scope_mapping_rows(spec),
        [2900, 6460],
    )
    add_paragraph(
        doc,
        "Kapsam dışında bırakılan noktalar bilinçli sınırlardır. Amaç üretim ortamı "
        "ölçeğinde gerçek şehir altyapısı kurmak değil, ilgili bulut servislerini doğru "
        "mimari noktaya yerleştirerek çalışan ve anlatılabilir bir mikroservis sistemi "
        "sunmaktır.",
    )

    doc.add_heading("3. Mikroservis sınırı ve sistem mimarisi", level=1)
    add_paragraph(doc, f"{spec.module_name} aşağıdaki sorumlulukları üstlenir:")
    add_bullets(doc, spec.service_boundary_do)
    add_paragraph(doc, "Bu servis özellikle şu görevleri üstlenmez:")
    add_bullets(doc, spec.service_boundary_dont)
    add_paragraph(
        doc,
        "Ortak platformda frontend yalnızca REST API çağrıları yapar. Servisler birbirinin "
        "DynamoDB tablolarına doğrudan bağlanmaz. Bu yaklaşım mikroservis veri sahipliği "
        "ilkesini korur ve deployment ortamında IAM yetkilerinin dar tutulmasını sağlar.",
    )
    diagram_path = ensure_architecture_diagram(spec)
    doc.add_paragraph()
    doc.add_picture(str(diagram_path), width=Inches(6.3))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"Şekil 1. {spec.module_name} mikroservis mimari diyagramı.")
    set_run_font(r, size=9, color=MUTED, italic=True)

    doc.add_heading("4. Teknoloji kararları", level=1)
    add_docx_table(
        doc,
        ["Bileşen", "Karar ve gerekçe"],
        [[a, b] for a, b in spec.technologies],
        [2300, 7060],
    )
    add_paragraph(
        doc,
        "Backend tarafında FastAPI seçimi, hızlı REST API geliştirme, otomatik Swagger "
        "dokümantasyonu ve Python ekosistemiyle uyum nedeniyle yapılmıştır. Frontend "
        "tarafında React + Vite + TypeScript kullanılmış; grafik ve canlı görünüm için "
        "TanStack Query polling, Recharts ve Leaflet tercih edilmiştir.",
    )

    doc.add_heading("5. Veri modeli ve API tasarımı", level=1)
    for table_name, fields in spec.data_tables:
        doc.add_heading(table_name, level=2)
        add_docx_table(doc, ["Alan", "Açıklama"], [[a, b] for a, b in fields], [2600, 6760])
    doc.add_heading("API yüzeyi", level=2)
    add_bullets(doc, spec.endpoints)
    add_paragraph(
        doc,
        "Endpoint listesi frontend'in ihtiyaç duyduğu okuma, durum ve demo kontrol akışlarını "
        "karşılar. Veri üretimi, her modülün kendi doğasına uygun giriş kanalından yapılır; "
        "örneğin IoT tarafında telemetry REST ile değil MQTT üzerinden sisteme girer.",
    )
    doc.add_heading("Local-AWS karşılıkları", level=2)
    add_docx_table(
        doc,
        ["Local bileşen", "AWS karşılığı"],
        local_aws_rows(spec),
        [3100, 6260],
    )

    doc.add_heading("6. Local çalışma ve demo akışı", level=1)
    add_paragraph(
        doc,
        "Local demo Docker Compose ile çalışacak şekilde tasarlanmıştır. Frontend tek paneldir; "
        "IoT, Video ve ML servisleri ayrı container süreçleri olarak ayağa kalkar. DynamoDB Local "
        "local veritabanı karşılığıdır ve Mosquitto yalnızca local IoT demo broker'ı olarak yer alır.",
    )
    add_docx_table(
        doc,
        ["Akış adımı", "Açıklama"],
        flow_rows(spec.local_flow),
        [2200, 7160],
    )
    if (ASSET_DIR / spec.screenshot).exists():
        doc.add_paragraph()
        doc.add_picture(str(ASSET_DIR / spec.screenshot), width=Inches(6.1))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Şekil 2. {spec.module_name} canlı deployment ekran görüntüsü.")
        set_run_font(r, size=9, color=MUTED, italic=True)

    doc.add_heading("7. AWS deployment mimarisi", level=1)
    add_paragraph(
        doc,
        "AWS ortamında servis, ayrı ECR repository'si, ayrı ECS Fargate service'i, ayrı task "
        "definition ve ayrı IAM task role ile konumlandırılmıştır. Frontend S3 + CloudFront "
        "üzerinden yayınlanır; `/api/*` path davranışı ALB'ye yönlendirilir. ALB path routing "
        "ile ilgili istek ilgili mikroservis target group'una aktarılır.",
    )
    add_docx_table(
        doc,
        ["AWS akış adımı", "Açıklama"],
        flow_rows(spec.aws_flow),
        [2200, 7160],
    )
    add_docx_table(
        doc,
        ["AWS kaynağı", "Bu rapordaki görevi"],
        [[a, b] for a, b in spec.deployed_resources],
        [3000, 6360],
    )
    doc.add_heading("Operasyon ve maliyet notları", level=2)
    add_docx_table(
        doc,
        ["Konu", "Not"],
        cost_rows(),
        [2600, 6760],
    )

    doc.add_heading("8. Güvenlik, IAM ve veri sahipliği", level=1)
    add_paragraph(
        doc,
        "Her mikroservise yalnızca kendi verisine erişim yetkisi verilmiştir. ECS task execution "
        "role sadece ECR image çekme ve CloudWatch Logs yazma gibi altyapı görevleri için; "
        "uygulama task role ise servis özelindeki DynamoDB/S3/Lambda erişimleri için kullanılır.",
    )
    add_paragraph(
        doc,
        "Bu ayrım, mikroservis sınırını hem kod hem deployment seviyesinde görünür hale getirir. "
        "Bir servisin diğer modülün tablosuna doğrudan erişmemesi, hata izolasyonu ve ilerde "
        "ölçekleme açısından da daha sağlıklı bir tasarımdır.",
    )

    doc.add_heading("9. Test, smoke doğrulama ve kanıtlar", level=1)
    add_docx_table(doc, ["Kontrol", "Sonuç"], [[a, b] for a, b in spec.validation], [3200, 6160])
    add_paragraph(
        doc,
        f"Canlı frontend URL'si {PUBLIC_URL} adresidir. Smoke testlerde CloudFront üzerinden "
        "dashboard ve ilgili modül ekranı açılmış, API hatası veya console error gözlenmemiştir.",
    )

    doc.add_heading("10. Ortak platform entegrasyonu", level=1)
    add_paragraph(
        doc,
        "Bu mikroservis tek başına anlamlı bir modüldür; ancak platform değeri frontend'de diğer "
        "servislerle birlikte görünür. IoT verileri şehir durumu ve harita kartlarını besler; "
        "video analizleri kamera yoğunluğu ve güvenlik göstergelerini üretir; ML servisi ise "
        "bu iki kaynaktan gelen veriyi risk ve öneriye dönüştürür.",
    )
    add_paragraph(
        doc,
        "Ortak frontend'in veritabanı yoktur. Bu bilinçli kararla frontend yalnızca görselleştirme "
        "ve kullanıcı etkileşimi katmanı olarak kalır. API Gateway kullanılmamış, bunun yerine "
        "CloudFront + ALB path routing ile sade ve gösterilebilir bir deployment kurgusu kurulmuştur.",
    )

    doc.add_heading("11. Karşılaşılan sorunlar ve çözüm kararları", level=1)
    add_docx_table(doc, ["Sorun / karar alanı", "Çözüm"], [[a, b] for a, b in spec.challenges], [3000, 6360])

    doc.add_heading("12. Sonuç", level=1)
    add_paragraph(doc, spec.conclusion)

    doc.add_heading("Kaynaklar", level=1)
    add_paragraph(doc, source_summary(spec))


def build_docx(spec: ReportSpec):
    doc = Document()
    setup_docx_styles(doc)
    set_header_footer(doc, spec)
    add_cover(doc, spec)
    add_front_matter(doc, spec)
    add_module_sections(doc, spec)
    out = OUT_DIR / f"{spec.slug}.docx"
    doc.save(out)
    return out


def register_pdf_fonts():
    base = Path("/usr/share/fonts/truetype/noto")
    pdfmetrics.registerFont(TTFont("NotoSans", str(base / "NotoSans-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("NotoSans-Bold", str(base / "NotoSans-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("NotoSans-Italic", str(base / "NotoSans-Italic.ttf")))
    pdfmetrics.registerFont(TTFont("NotoSans-BoldItalic", str(base / "NotoSans-BoldItalic.ttf")))
    pdfmetrics.registerFont(TTFont("NotoSansMono", str(base / "NotoSansMono-Regular.ttf")))


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title",
            parent=styles["Title"],
            fontName="NotoSans-Bold",
            fontSize=22,
            leading=27,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=10,
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            parent=styles["Normal"],
            fontName="NotoSans-Italic",
            fontSize=12,
            leading=16,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#5B6776"),
            spaceAfter=18,
        ),
        "h1": ParagraphStyle(
            "h1",
            parent=styles["Heading1"],
            fontName="NotoSans-Bold",
            fontSize=15,
            leading=18,
            textColor=colors.HexColor("#2E74B5"),
            spaceBefore=10,
            spaceAfter=6,
        ),
        "h2": ParagraphStyle(
            "h2",
            parent=styles["Heading2"],
            fontName="NotoSans-Bold",
            fontSize=12.5,
            leading=15,
            textColor=colors.HexColor("#2E74B5"),
            spaceBefore=8,
            spaceAfter=5,
        ),
        "h2_center": ParagraphStyle(
            "h2_center",
            parent=styles["Heading2"],
            fontName="NotoSans-Bold",
            fontSize=11,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#2E74B5"),
            spaceBefore=3,
            spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "body",
            parent=styles["BodyText"],
            fontName="NotoSans",
            fontSize=9.6,
            leading=12.5,
            alignment=TA_JUSTIFY,
            spaceAfter=6,
        ),
        "body_left": ParagraphStyle(
            "body_left",
            parent=styles["BodyText"],
            fontName="NotoSans",
            fontSize=9.6,
            leading=12.5,
            alignment=TA_LEFT,
            spaceAfter=6,
        ),
        "small": ParagraphStyle(
            "small",
            parent=styles["BodyText"],
            fontName="NotoSans",
            fontSize=8.2,
            leading=10.5,
            textColor=colors.HexColor("#475569"),
            spaceAfter=4,
        ),
        "caption": ParagraphStyle(
            "caption",
            parent=styles["BodyText"],
            fontName="NotoSans-Italic",
            fontSize=8.4,
            leading=10.5,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#5B6776"),
            spaceBefore=4,
            spaceAfter=8,
        ),
        "code": ParagraphStyle(
            "code",
            parent=styles["Code"],
            fontName="NotoSansMono",
            fontSize=8.2,
            leading=10,
            textColor=colors.HexColor("#0F172A"),
            backColor=colors.HexColor("#F8FAFC"),
            borderColor=colors.HexColor("#E2E8F0"),
            borderWidth=0.25,
            borderPadding=6,
            spaceAfter=8,
        ),
    }


def p(text: str, style):
    return Paragraph(escape(text).replace("\n", "<br/>"), style)


def pdf_table(headers: list[str], rows: list[list[str]], widths: list[float], styles_map):
    data = [[p(h, styles_map["small"]) for h in headers]]
    for row in rows:
        data.append([p(str(v), styles_map["small"]) for v in row])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="CENTER")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("FONTNAME", (0, 0), (-1, 0), "NotoSans-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DADCE0")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def bullet_flow(items: list[str], styles_map):
    return ListFlowable(
        [ListItem(p(item, styles_map["body_left"]), leftIndent=12) for item in items],
        bulletType="bullet",
        leftIndent=18,
        bulletFontName="NotoSans",
        bulletFontSize=8,
        spaceAfter=6,
    )


def numbered_flow(items: list[str], styles_map):
    return ListFlowable(
        [ListItem(p(item, styles_map["body_left"]), leftIndent=12) for item in items],
        bulletType="1",
        leftIndent=18,
        bulletFontName="NotoSans",
        bulletFontSize=8,
        spaceAfter=6,
    )


def add_pdf_section(story: list, title: str, styles_map):
    story.append(p(title, styles_map["h1"]))


def add_pdf_h2(story: list, title: str, styles_map):
    story.append(p(title, styles_map["h2"]))


def on_page(canvas, doc, spec: ReportSpec):
    canvas.saveState()
    if doc.page > 1:
        canvas.setFont("NotoSans", 8)
        canvas.setFillColor(colors.HexColor("#5B6776"))
        canvas.drawString(72, 740, f"{PROJECT} | {spec.module_name}")
        canvas.drawRightString(540, 42, f"Sayfa {doc.page}")
    canvas.restoreState()


def build_pdf(spec: ReportSpec):
    register_pdf_fonts()
    styles_map = pdf_styles()
    out = OUT_DIR / f"{spec.slug}.pdf"
    doc = SimpleDocTemplate(
        str(out),
        pagesize=letter,
        rightMargin=72,
        leftMargin=72,
        topMargin=60,
        bottomMargin=48,
        title=spec.title,
        author=STUDENT,
    )
    story: list = []
    story.append(Spacer(1, 0.36 * inch))
    story.append(p(COURSE, styles_map["subtitle"]))
    story.append(p(PROJECT, styles_map["title"]))
    story.append(p(spec.title, styles_map["h1"]))
    story.append(p(spec.scope, styles_map["subtitle"]))
    metadata = [
        ["Öğrenci", STUDENT],
        ["Öğrenci No", STUDENT_NO],
        ["Ders", COURSE],
        ["Öğretim Elemanı", INSTRUCTOR],
        ["Frontend Repo Linki", FRONTEND_REPO],
        [spec.repo_label, spec.repo_url],
        ["Demo Video Linki", spec.demo_video_url],
    ]
    story.append(Spacer(1, 0.25 * inch))
    story.append(pdf_table(["Alan", "Değer"], metadata, [1.7 * inch, 4.55 * inch], styles_map))
    story.append(Spacer(1, 0.16 * inch))
    story.append(p("Kapak proje özeti", styles_map["h2_center"]))
    story.append(
        pdf_table(
            ["Alan", "Özet"],
            [[label, value] for label, value in cover_summary_rows(spec)],
            [1.7 * inch, 4.55 * inch],
            styles_map,
        )
    )
    story.append(PageBreak())

    add_pdf_section(story, "Özet", styles_map)
    story.append(p(spec.abstract, styles_map["body"]))
    story.append(p("Anahtar Kelimeler: " + ", ".join(spec.keywords), styles_map["body_left"]))
    add_pdf_section(story, "İçindekiler", styles_map)
    story.append(
        numbered_flow(
            [
                "Giriş",
                "Kapsam",
                "Mikroservis sınırı ve sistem mimarisi",
                "Teknoloji kararları",
                "Veri modeli ve API tasarımı",
                "Local çalışma ve demo akışı",
                "AWS deployment mimarisi",
                "Güvenlik, IAM ve veri sahipliği",
                "Test, smoke doğrulama ve kanıtlar",
                "Ortak platform entegrasyonu",
                "Karşılaşılan sorunlar ve çözüm kararları",
                "Sonuç",
                "Kaynaklar",
            ],
            styles_map,
        )
    )
    story.append(PageBreak())

    add_pdf_section(story, "1. Giriş", styles_map)
    story.append(
        p(
            f"Bu rapor, {PROJECT} kapsamında geliştirilen {spec.module_name} bileşenini "
            "ayrı bir mikroservis olarak ele alır. Platform tek bir yönetim paneli gibi "
            "kullanılsa da her mikroservis kendi veri modeline, kendi API yüzeyine, kendi "
            "container image'ına ve kendi AWS yetki sınırına sahiptir.",
            styles_map["body"],
        )
    )
    story.append(
        p(
            "Temel tasarım hedefi, ders kapsamındaki bulut servislerini yalnızca isim olarak "
            "kullanmak değil, doğru problemi doğru AWS servisi ile eşleştirmektir. Bu nedenle "
            "raporda local demo akışı ile AWS deployment karşılığı birlikte anlatılmıştır.",
            styles_map["body"],
        )
    )

    add_pdf_section(story, "2. Kapsam", styles_map)
    story.append(
        p(
            f"Bu bölüm, {spec.module_name} mikroservisinin proje içindeki sınırını ve "
            f"{COURSE} final yönergesindeki {spec.pdf_project} beklentilerini nasıl "
            "karşıladığını özetler. Kapsam, yalnızca uygulama ekranını değil; veri giriş "
            "kanalını, servis sorumluluğunu, AWS karşılığını ve raporda savunulacak deployment "
            "kararlarını birlikte ele alır.",
            styles_map["body"],
        )
    )
    story.append(
        pdf_table(
            ["Kapsam maddesi", "Projede karşılığı"],
            scope_mapping_rows(spec),
            [1.8 * inch, 4.45 * inch],
            styles_map,
        )
    )
    story.append(
        p(
            "Kapsam dışında bırakılan noktalar bilinçli sınırlardır. Amaç üretim ortamı "
            "ölçeğinde gerçek şehir altyapısı kurmak değil, ilgili bulut servislerini doğru "
            "mimari noktaya yerleştirerek çalışan ve anlatılabilir bir mikroservis sistemi sunmaktır.",
            styles_map["body"],
        )
    )

    add_pdf_section(story, "3. Mikroservis sınırı ve sistem mimarisi", styles_map)
    story.append(p(f"{spec.module_name} aşağıdaki sorumlulukları üstlenir:", styles_map["body_left"]))
    story.append(bullet_flow(spec.service_boundary_do, styles_map))
    story.append(p("Bu servis özellikle şu görevleri üstlenmez:", styles_map["body_left"]))
    story.append(bullet_flow(spec.service_boundary_dont, styles_map))
    story.append(
        p(
            "Ortak platformda frontend yalnızca REST API çağrıları yapar. Servisler birbirinin "
            "DynamoDB tablolarına doğrudan bağlanmaz. Bu yaklaşım mikroservis veri sahipliği "
            "ilkesini korur ve deployment ortamında IAM yetkilerinin dar tutulmasını sağlar.",
            styles_map["body"],
        )
    )
    diagram_path = ensure_architecture_diagram(spec)
    diagram = Image(str(diagram_path))
    diagram._restrictSize(6.25 * inch, 2.9 * inch)
    story.append(KeepTogether([diagram, p(f"Şekil 1. {spec.module_name} mikroservis mimari diyagramı.", styles_map["caption"])]))

    add_pdf_section(story, "4. Teknoloji kararları", styles_map)
    story.append(pdf_table(["Bileşen", "Karar ve gerekçe"], [[a, b] for a, b in spec.technologies], [1.55 * inch, 4.7 * inch], styles_map))
    story.append(
        p(
            "Backend tarafında FastAPI seçimi, hızlı REST API geliştirme, otomatik Swagger "
            "dokümantasyonu ve Python ekosistemiyle uyum nedeniyle yapılmıştır. Frontend tarafında "
            "React + Vite + TypeScript kullanılmış; canlı görünüm için TanStack Query polling tercih edilmiştir.",
            styles_map["body"],
        )
    )

    add_pdf_section(story, "5. Veri modeli ve API tasarımı", styles_map)
    for table_name, fields in spec.data_tables:
        add_pdf_h2(story, table_name, styles_map)
        story.append(pdf_table(["Alan", "Açıklama"], [[a, b] for a, b in fields], [1.75 * inch, 4.5 * inch], styles_map))
        story.append(Spacer(1, 0.08 * inch))
    add_pdf_h2(story, "API yüzeyi", styles_map)
    story.append(bullet_flow(spec.endpoints, styles_map))
    story.append(
        p(
            "Endpoint listesi frontend'in ihtiyaç duyduğu okuma, durum ve demo kontrol akışlarını "
            "karşılar. Veri üretimi, her modülün kendi doğasına uygun giriş kanalından yapılır.",
            styles_map["body"],
        )
    )
    add_pdf_h2(story, "Local-AWS karşılıkları", styles_map)
    story.append(
        pdf_table(
            ["Local bileşen", "AWS karşılığı"],
            local_aws_rows(spec),
            [2.05 * inch, 4.2 * inch],
            styles_map,
        )
    )
    add_pdf_section(story, "6. Local çalışma ve demo akışı", styles_map)
    story.append(
        p(
            "Local demo Docker Compose ile çalışacak şekilde tasarlanmıştır. Frontend tek paneldir; "
            "IoT, Video ve ML servisleri ayrı container süreçleri olarak ayağa kalkar. DynamoDB Local "
            "local veritabanı karşılığıdır.",
            styles_map["body"],
        )
    )
    story.append(pdf_table(["Akış adımı", "Açıklama"], flow_rows(spec.local_flow), [1.35 * inch, 4.9 * inch], styles_map))
    img_path = ASSET_DIR / spec.screenshot
    if img_path.exists():
        img = Image(str(img_path))
        img._restrictSize(6.25 * inch, 3.05 * inch)
        story.append(KeepTogether([img, p(f"Şekil 2. {spec.module_name} canlı deployment ekran görüntüsü.", styles_map["caption"])]))

    add_pdf_section(story, "7. AWS deployment mimarisi", styles_map)
    story.append(
        p(
            "AWS ortamında servis, ayrı ECR repository'si, ayrı ECS Fargate service'i, ayrı task "
            "definition ve ayrı IAM task role ile konumlandırılmıştır. Frontend S3 + CloudFront "
            "üzerinden yayınlanır; `/api/*` path davranışı ALB'ye yönlendirilir.",
            styles_map["body"],
        )
    )
    story.append(pdf_table(["AWS akış adımı", "Açıklama"], flow_rows(spec.aws_flow), [1.35 * inch, 4.9 * inch], styles_map))
    story.append(pdf_table(["AWS kaynağı", "Bu rapordaki görevi"], [[a, b] for a, b in spec.deployed_resources], [2.0 * inch, 4.25 * inch], styles_map))
    story.append(
        p(
            "Bu bölümdeki kaynaklar, local Docker Compose mimarisinin buluttaki karşılığıdır. "
            "Amaç backendleri yalnızca bir sunucuya koymak değil; veri girişini, analiz işini, "
            "saklama katmanını ve API yüzeyini doğru managed AWS servisleriyle ayırmaktır.",
            styles_map["body"],
        )
    )
    add_pdf_h2(story, "Operasyon ve maliyet notları", styles_map)
    story.append(
        pdf_table(
            ["Konu", "Not"],
            cost_rows(),
            [1.8 * inch, 4.45 * inch],
            styles_map,
        )
    )

    add_pdf_section(story, "8. Güvenlik, IAM ve veri sahipliği", styles_map)
    story.append(
        p(
            "Her mikroservise yalnızca kendi verisine erişim yetkisi verilmiştir. ECS task execution "
            "role sadece ECR image çekme ve CloudWatch Logs yazma gibi altyapı görevleri için; "
            "uygulama task role ise servis özelindeki DynamoDB/S3/Lambda erişimleri için kullanılır.",
            styles_map["body"],
        )
    )
    story.append(
        p(
            "Bu ayrım, mikroservis sınırını hem kod hem deployment seviyesinde görünür hale getirir. "
            "Bir servisin diğer modülün tablosuna doğrudan erişmemesi, hata izolasyonu ve ilerde "
            "ölçekleme açısından daha sağlıklı bir tasarımdır.",
            styles_map["body"],
        )
    )

    add_pdf_section(story, "9. Test, smoke doğrulama ve kanıtlar", styles_map)
    story.append(pdf_table(["Kontrol", "Sonuç"], [[a, b] for a, b in spec.validation], [2.05 * inch, 4.2 * inch], styles_map))
    story.append(p(f"Canlı frontend URL'si {PUBLIC_URL} adresidir.", styles_map["body_left"]))
    story.append(
        p(
            "Bu doğrulamalar teslim videosunda anlatılacak akış için de kanıt niteliğindedir: "
            "önce canlı frontend açılır, ardından modül ekranı gösterilir, en son API ve AWS "
            "kaynak eşleşmesi açıklanır.",
            styles_map["body"],
        )
    )

    add_pdf_section(story, "10. Ortak platform entegrasyonu", styles_map)
    story.append(
        p(
            "Bu mikroservis tek başına anlamlı bir modüldür; ancak platform değeri frontend'de diğer "
            "servislerle birlikte görünür. IoT verileri şehir durumunu, video analizleri kamera "
            "yoğunluğunu, ML servisi ise risk ve öneri katmanını üretir.",
            styles_map["body"],
        )
    )
    story.append(
        p(
            "Ortak frontend'in veritabanı yoktur. API Gateway kullanılmamış, bunun yerine "
            "CloudFront + ALB path routing ile sade ve gösterilebilir bir deployment kurgusu kurulmuştur.",
            styles_map["body"],
        )
    )

    add_pdf_section(story, "11. Karşılaşılan sorunlar ve çözüm kararları", styles_map)
    story.append(pdf_table(["Sorun / karar alanı", "Çözüm"], [[a, b] for a, b in spec.challenges], [2.0 * inch, 4.25 * inch], styles_map))

    add_pdf_section(story, "12. Sonuç", styles_map)
    story.append(p(spec.conclusion, styles_map["body"]))

    add_pdf_section(story, "Kaynaklar", styles_map)
    story.append(p(source_summary(spec), styles_map["small"]))
    doc.build(story, onFirstPage=lambda c, d: on_page(c, d, spec), onLaterPages=lambda c, d: on_page(c, d, spec))
    return out


def specs() -> list[ReportSpec]:
    common_sources = [
        "3522 Bulut Bilişim dersi final proje yönergesi.",
        "AWS dokümantasyonu: ECS Fargate, DynamoDB, S3, CloudFront, Lambda ve ilgili managed servisler.",
        "FastAPI dokümantasyonu.",
        "React, Vite, TanStack Query ve Recharts dokümantasyonları.",
        "Proje repository dokümanları: architecture-decisions, deployment-plan ve aws-deployment-runbook.",
    ]
    return [
        ReportSpec(
            slug="iot-akilli-sehir-raporu",
            title="IoT ve Akıllı Şehir Telemetri Mikroservisi",
            scope="Akıllı şehir sensör verisi toplama, MQTT akışı, DynamoDB event saklama ve AWS IoT Core deployment kapsamı",
            module_name="IoT Service",
            repo_label="IoT Service Repo Linki",
            repo_url="https://github.com/eemrezcan/bulutproje3-iot-service",
            demo_video_url="https://youtu.be/cwmFWNG1RY0",
            screenshot="iot-readings.jpg",
            abstract=(
                "Bu rapor, Akıllı Şehir Yönetim Platformu içinde geliştirilen IoT Service mikroservisini inceler. "
                "Servis; şehir bölgelerinden sıcaklık, nem, hava kalitesi ve trafik yoğunluğu gibi telemetry "
                "verilerini MQTT üzerinden alır, veriyi DynamoDB tablosunda zaman serisi event olarak saklar ve "
                "frontend'e REST API ile sunar. Local ortamda Mosquitto ve DynamoDB Local kullanılırken AWS "
                "ortamında AWS IoT Core, IoT Rule, Lambda, DynamoDB, ECS Fargate, ALB ve CloudFront bileşenleri "
                "ile karşılık kurulmuştur."
            ),
            keywords=["IoT", "MQTT", "AWS IoT Core", "DynamoDB", "FastAPI", "ECS Fargate", "Akıllı Şehir"],
            pdf_project="Proje 7 - IoT ve Akıllı Şehir Uygulaması",
            pdf_expectations=[
                "IoT cihazları üzerinden veri toplama ve raporlama yapılması.",
                "MQTT veya benzeri IoT protokolünün kullanılması.",
                "AWS IoT Core ve Lambda ile bulut tabanlı veri işleme akışının kurulması.",
                "Toplanan verinin gerçek zamanlıya yakın görselleştirilmesi.",
            ],
            service_boundary_do=[
                "Sensör cihazlarını, bölgeleri ve son telemetry durumlarını yönetir.",
                "MQTT topic'lerinden gelen telemetry mesajlarını işler.",
                "Sıcaklık, nem, hava kalitesi ve trafik yoğunluğu ölçümlerini saklar.",
                "Frontend için sensör listesi, son ölçüm ve bölge geçmişi API'leri sunar.",
            ],
            service_boundary_dont=[
                "Kamera veya video akışı işlemez.",
                "Makine öğrenmesi tahmini üretmez.",
                "Video veya ML servislerinin veritabanlarına doğrudan bağlanmaz.",
            ],
            technologies=[
                ("Python + FastAPI", "REST API, health check ve sensör okuma endpointleri için kullanıldı."),
                ("MQTT", "Telemetry girişinin gerçek IoT publish/subscribe modeline benzemesi için seçildi."),
                ("Mosquitto", "Local Docker Compose demosunda AWS IoT Core'un yerel karşılığıdır."),
                ("DynamoDB", "Sensör metadata ve zaman damgalı ölçüm eventleri için kullanıldı."),
                ("AWS IoT Core + IoT Rule + Lambda", "AWS tarafında MQTT mesajını alıp DynamoDB'ye yazan bulut hattıdır."),
            ],
            local_flow=(
                "Sensor Simulator\n"
                "  -> MQTT publish\n"
                "  -> Mosquitto Broker\n"
                "  -> IoT Service MQTT Consumer\n"
                "  -> DynamoDB Local\n"
                "  -> IoT Service REST API\n"
                "  -> Frontend IoT ekranı"
            ),
            aws_flow=(
                "Sensor / simulator\n"
                "  -> MQTT\n"
                "  -> AWS IoT Core\n"
                "  -> IoT Rule\n"
                "  -> Lambda: smart-city-final-iot-ingest\n"
                "  -> DynamoDB: iot_sensor_readings\n"
                "  -> ECS IoT API\n"
                "  -> CloudFront frontend"
            ),
            data_tables=[
                (
                    "iot_sensors",
                    [
                        ("sensor_id", "Partition key; sensör kimliği."),
                        ("zone", "Meydan, Otogar, Kampus, Hastane veya Sanayi bölgesi."),
                        ("name", "Kullanıcı arayüzünde görünen sensör adı."),
                        ("status", "active, passive veya maintenance durumu."),
                        ("latitude / longitude", "Harita gösterimi için konum bilgisi."),
                    ],
                ),
                (
                    "iot_sensor_readings",
                    [
                        ("sensor_id", "Partition key; ölçümün ait olduğu sensör."),
                        ("timestamp", "Sort key; zaman sıralı telemetry saklama."),
                        ("temperature / humidity", "Çevresel sensör değerleri."),
                        ("air_quality_index", "Hava kalitesi endeksi."),
                        ("traffic_level", "Trafik yoğunluğu göstergesi."),
                        ("status_level", "normal, warning veya critical sınıflandırması."),
                    ],
                ),
            ],
            endpoints=[
                "GET /health",
                "GET /sensors",
                "GET /readings/latest",
                "GET /readings?sensor_id={sensor_id}&limit=50",
                "GET /zones/{zone}/readings?limit=50",
                "POST /simulator/start",
                "POST /simulator/stop",
            ],
            deployed_resources=[
                ("ECS Fargate service", "smart-city-final-iot-service API container'ını çalıştırır."),
                ("ECR repository", "smart-city-final/iot-service image versiyonlarını saklar."),
                ("DynamoDB", "iot_sensors ve iot_sensor_readings tablolarını barındırır."),
                ("AWS IoT Core", "smart-city/iot/+/+/telemetry topic'lerini MQTT giriş kanalı olarak alır."),
                ("IoT Rule + Lambda", "Telemetry payload'ini doğrular ve DynamoDB'ye yazar."),
                ("CloudFront + ALB", "/api/iot path'ini IoT target group'una yönlendirir."),
            ],
            validation=[
                ("CloudFront IoT API health", "/api/iot/health status=ok döndürdü."),
                ("IoT readings", "/api/iot/readings/latest beş bölge için son ölçümleri döndürdü."),
                ("IoT Core test", "smart-city/iot/meydan/sensor_meydan_01/telemetry topic'ine publish edilen mesaj DynamoDB'ye yazıldı."),
                ("Frontend smoke", "IoT ekranında Hastane, Kampus, Meydan, Otogar ve Sanayi kartları göründü."),
                ("ECS service", "desired=1, running=1, target health=healthy."),
            ],
            challenges=[
                ("Local ve AWS MQTT ayrımı", "Localde Mosquitto kullanıldı; AWS'te MQTT hattı IoT Core + Rule + Lambda ile temsil edildi."),
                ("Servis verisi sahipliği", "ML veya Video servisleri IoT tablolarına doğrudan bağlanmadı; veri REST API ile paylaşıldı."),
                ("Demo kararlılığı", "Sentetik sensör üretimi kontrollü değer aralıklarıyla yapıldı ve beş bölge için stabil veri sağlandı."),
            ],
            conclusion=(
                "IoT Service, Proje 7 kapsamındaki MQTT tabanlı cihaz haberleşmesi, bulut veri işleme ve gerçek zamanlıya yakın "
                "görselleştirme beklentilerini karşılar. AWS tarafında IoT Core + Rule + Lambda + DynamoDB hattı kurulmuş, "
                "API katmanı ise ayrı ECS Fargate servisi olarak deploy edilmiştir."
            ),
            sources=common_sources + ["AWS IoT Core ve Eclipse Mosquitto dokümantasyonları."],
        ),
        ReportSpec(
            slug="video-akilli-sehir-raporu",
            title="Video Akışı ve Akıllı Şehir Kamera Analizi Mikroservisi",
            scope="Kamera akışı, canlı şehir sahnesi, analiz eventleri, Rekognition stored video hattı ve DynamoDB event store kapsamı",
            module_name="Video Analysis Service",
            repo_label="Video Analysis Service Repo Linki",
            repo_url="https://github.com/eemrezcan/bulutproje4-video-analysis-service",
            demo_video_url="https://youtu.be/5lllkAunq5A",
            screenshot="video-live-feed.jpg",
            abstract=(
                "Bu rapor, Akıllı Şehir Yönetim Platformu içindeki Video Analysis Service mikroservisini ele alır. "
                "Servis şehir kameralarını listeler, canlı kamera hissi veren MJPEG akışları sunar ve kişi sayısı, "
                "araç sayısı, hareket seviyesi, kalabalık seviyesi ve recognition label gibi analiz eventlerini üretir. "
                "AWS tarafında video dosyası S3'te tutulmuş, Rekognition Video stored analysis SNS ve Lambda üzerinden "
                "DynamoDB'ye bağlanmıştır. Video dosyaları veritabanında değil object storage üzerinde tutulur."
            ),
            keywords=["Video Akışı", "Rekognition Video", "S3", "SNS", "Lambda", "DynamoDB", "ECS Fargate"],
            pdf_project="Proje 6 - Video Akışı ve İşleme Uygulaması",
            pdf_expectations=[
                "Video akışlarını yönetme ve kullanıcıya oynatma/görüntüleme imkanı sunma.",
                "RTMP, WebRTC veya Kinesis Video Streams gibi video akışı yaklaşımını tartışma.",
                "AWS Rekognition gibi bulut tabanlı video işleme servisiyle analiz üretme.",
                "Nesne tanıma, etiketleme ve analiz sonuçlarını saklama.",
            ],
            service_boundary_do=[
                "Kamera metadata'sını ve şehir bölgelerini yönetir.",
                "Canlı kamera hissi veren MJPEG akış endpointleri sunar.",
                "Kişi sayısı, araç sayısı, hareket ve kalabalık seviyesi eventleri üretir.",
                "Video analiz sonuçlarını kendi DynamoDB tablosunda saklar.",
            ],
            service_boundary_dont=[
                "IoT sensör verisi toplamaz.",
                "Sıcaklık veya hava kalitesi üretmez.",
                "ML tahmini yapmaz ve ML modelini eğitmez.",
                "Video dosyasını DynamoDB'de saklamaz.",
            ],
            technologies=[
                ("Python + FastAPI", "Kamera, stream ve analiz endpointleri için kullanıldı."),
                ("OpenCV / sentetik sahne", "Local ve ECS demosunda stabil canlı şehir görüntüsü üretmek için kullanıldı."),
                ("MJPEG live endpoint", "Browser tarafında canlı akış hissi için live.mjpg görsel akış endpoint'i uygulandı."),
                ("Amazon S3", "Video kaynak ve Rekognition input dosyalarını saklar."),
                ("Rekognition Video + SNS + Lambda", "Stored video analysis sonucunu asenkron işleyip DynamoDB'ye yazar."),
                ("DynamoDB", "Kamera metadata ve zaman damgalı analiz eventleri için kullanılır."),
            ],
            local_flow=(
                "Synthetic city scene / video loop\n"
                "  -> Video Service live MJPEG endpoint\n"
                "  -> Frontend camera panel\n"
                "  -> Video analyzer\n"
                "  -> DynamoDB Local\n"
                "  -> Video Service REST API\n"
                "  -> ML ve Dashboard ekranları"
            ),
            aws_flow=(
                "S3 video object\n"
                "  -> Rekognition Video stored analysis\n"
                "  -> SNS notification\n"
                "  -> Lambda: smart-city-final-video-result\n"
                "  -> DynamoDB: video_analysis_results\n"
                "  -> ECS Video API\n"
                "  -> CloudFront frontend\n\n"
                "Canlı demo görüntüsü:\n"
                "ECS Video API -> /api/video/cameras/{id}/live.mjpg -> Frontend"
            ),
            data_tables=[
                (
                    "video_cameras",
                    [
                        ("camera_id", "Partition key; kamera kimliği."),
                        ("zone", "Kameranın bulunduğu şehir bölgesi."),
                        ("name", "Frontend kamera kartında görünen ad."),
                        ("stream_url", "Canlı görüntü endpoint veya stream yolu."),
                        ("status", "active, passive veya maintenance."),
                        ("latitude / longitude", "Harita entegrasyonu için konum."),
                    ],
                ),
                (
                    "video_analysis_results",
                    [
                        ("camera_id", "Partition key; analiz eventinin ait olduğu kamera."),
                        ("timestamp", "Sort key; zaman damgalı analiz geçmişi."),
                        ("zone", "Bölge bazlı sorgular için alan."),
                        ("people_count / vehicle_count", "Kişi ve araç sayısı."),
                        ("motion_level / crowd_level", "low, medium, high veya critical sınıflandırması."),
                        ("recognition_labels", "crowd, traffic, vehicle, emergency gibi etiketler."),
                    ],
                ),
            ],
            endpoints=[
                "GET /health",
                "GET /cameras",
                "GET /cameras/{camera_id}/stream",
                "GET /cameras/{camera_id}/live.mjpg",
                "GET /analysis/latest",
                "GET /analysis?camera_id={camera_id}&limit=50",
                "GET /zones/{zone}/analysis?limit=50",
                "POST /analysis/generate",
                "POST /analyzer/start",
                "POST /analyzer/stop",
            ],
            deployed_resources=[
                ("ECS Fargate service", "smart-city-final-video-service API ve canlı MJPEG endpointlerini çalıştırır."),
                ("ECR repository", "smart-city-final/video-service image versiyonlarını saklar."),
                ("S3 video bucket", "Rekognition input ve video kaynak dosyalarını barındırır."),
                ("Rekognition Video", "S3 üzerindeki H.264 MP4 için label detection çalıştırır."),
                ("SNS + Lambda", "Rekognition job tamamlanınca sonucu DynamoDB eventine dönüştürür."),
                ("DynamoDB", "video_cameras ve video_analysis_results tablolarını barındırır."),
                ("CloudFront + ALB", "/api/video path'ini Video target group'una yönlendirir."),
            ],
            validation=[
                ("CloudFront Video API health", "/api/video/health status=ok döndürdü."),
                ("Canlı stream metadata", "stream endpoint mode=live_synthetic_*_mjpeg ve media_type=multipart/x-mixed-replace döndürdü."),
                ("Frontend smoke", "Video ekranında beş kamera listelendi ve seçili kamera görseli 960x540 yüklendi."),
                ("Rekognition job", "Stored video analysis SUCCEEDED oldu ve Lambda sonucu DynamoDB'ye yazdı."),
                ("ECS service", "desired=1, running=1, target health=healthy."),
            ],
            challenges=[
                ("Demo görüntüsü", "Test pattern kaldırıldı; yol, araç, yaya ve CANLI overlay içeren sentetik şehir sahnesi üretildi."),
                ("Codec uyumu", "Browser uyumu için WebM fallback denendi; final canlı görünümde MJPEG endpoint tercih edildi."),
                ("Kinesis kararı", "Kinesis doğru AWS karşılığı olarak planlandı; bu deployda canlı demo ECS MJPEG, analiz S3 + Rekognition ile tamamlandı."),
                ("DynamoDB seçimi", "Video dosyası DB'de tutulmadı; DynamoDB yalnızca zaman damgalı analiz eventleri için kullanıldı."),
            ],
            conclusion=(
                "Video Analysis Service, video akışı ve bulut analizini akıllı şehir kamera senaryosunda gösterir. "
                "Canlı görüntü ECS MJPEG endpointinden; bulut analizi S3 + Rekognition + SNS + Lambda + DynamoDB "
                "hattından sağlanmıştır."
            ),
            sources=common_sources
            + ["Amazon Rekognition Video, SNS, Kinesis Video Streams ve OpenCV dokümantasyonları."],
        ),
        ReportSpec(
            slug="ml-akilli-sehir-raporu",
            title="Akıllı Veri Analitiği ve Risk Tahmini Mikroservisi",
            scope="IoT ve video verilerinden özellik çıkarımı, scikit-learn modeli, rule-based risk motoru ve Lambda inference kapsamı",
            module_name="ML Analytics Service",
            repo_label="ML Analytics Service Repo Linki",
            repo_url="https://github.com/eemrezcan/bulutproje5-ml-analytics-service",
            demo_video_url="https://youtu.be/RbHlxU_dOfk",
            screenshot="ml-risk-summary.jpg",
            abstract=(
                "Bu rapor, Akıllı Şehir Yönetim Platformu içindeki ML Analytics Service mikroservisini inceler. Servis, "
                "IoT ve Video servislerinin REST API'lerinden aldığı son verileri bölge bazında birleştirir; scikit-learn "
                "ve joblib tabanlı model artifact'i ile sıcaklık/kalabalık tahmini üretir; rule-based risk motoru ile "
                "risk skoru ve karar destek önerisi oluşturur. AWS deploymentta ML API ayrı ECS Fargate servisi olarak "
                "çalıştırılmış, düşük maliyetli inference için Lambda alternatifi uygulanmış ve tahminler DynamoDB'deki "
                "ml_predictions tablosuna yazılmıştır."
            ),
            keywords=["Makine Öğrenmesi", "Scikit-learn", "Pandas", "Joblib", "Lambda Inference", "DynamoDB", "Risk Skoru"],
            pdf_project="Proje 3 - Akıllı Veri Analitiği ve Makine Öğrenmesi Uygulaması",
            pdf_expectations=[
                "Makine öğrenmesi modeli geliştirme ve model çıktısını uygulama içinde kullanma.",
                "Bulut ortamında veri işleme, modelleme ve inference yaklaşımı sunma.",
                "Scikit-learn, TensorFlow veya PyTorch gibi bir ML kütüphanesi kullanma.",
                "AWS SageMaker veya Lambda gibi servislerle model dağıtımı kararını açıklama.",
                "Veriden bilgi çıkarma, tahmin ve karar destek sonucu üretme.",
            ],
            service_boundary_do=[
                "IoT Service API'sinden sıcaklık, nem, hava kalitesi ve trafik verisi alır.",
                "Video Service API'sinden kişi sayısı, araç sayısı ve kalabalık seviyesi alır.",
                "Bölge ve zaman bazlı feature extraction yapar.",
                "Sıcaklık, kalabalık seviyesi, risk skoru ve öneri üretir.",
                "Tahmin sonuçlarını kendi DynamoDB tablosunda saklar.",
            ],
            service_boundary_dont=[
                "Sensör verisi üretmez veya MQTT dinlemez.",
                "Kamera/video işlemez, OpenCV veya Rekognition çalıştırmaz.",
                "IoT veya Video servislerinin DynamoDB tablolarına doğrudan bağlanmaz.",
            ],
            technologies=[
                ("Python + FastAPI", "Model durumu, eğitim ve tahmin endpointleri için kullanıldı."),
                ("scikit-learn", "Sıcaklık ve kalabalık tahmini için gerçek ML kütüphanesi olarak kullanıldı."),
                ("Pandas + Joblib", "Veri hazırlama ve model artifact kaydetme/yükleme için kullanıldı."),
                ("Rule-based risk engine", "Risk skorunu açıklanabilir ve demo sırasında stabil tutmak için kullanıldı."),
                ("Lambda inference", "SageMaker maliyeti yerine düşük maliyetli AWS inference alternatifi olarak deploy edildi."),
                ("DynamoDB", "Bölge bazlı tahmin ve risk eventlerini saklar."),
            ],
            local_flow=(
                "ML Analytics Service\n"
                "  -> IoT Service REST API\n"
                "  -> Video Analysis Service REST API\n"
                "  -> Feature extraction\n"
                "  -> scikit-learn model + rule-based risk engine\n"
                "  -> DynamoDB Local: ml_predictions\n"
                "  -> ML Service REST API\n"
                "  -> Frontend ML ekranı"
            ),
            aws_flow=(
                "CloudFront / API path routing\n"
                "  -> ECS ML API\n"
                "  -> IoT API + Video API REST okuma\n"
                "  -> Feature extraction\n"
                "  -> Lambda: smart-city-final-ml-inference\n"
                "  -> DynamoDB: ml_predictions\n"
                "  -> Frontend ML Analytics ekranı\n\n"
                "Ölçekli alternatif: SageMaker endpoint"
            ),
            data_tables=[
                (
                    "ml_predictions",
                    [
                        ("zone", "Partition key; tahminin ait olduğu şehir bölgesi."),
                        ("prediction_timestamp", "Sort key; tahminin üretildiği zaman."),
                        ("prediction_for", "Tahminin hedef zamanı; bu projede 60 dakika ufku."),
                        ("predicted_temperature", "Bir saat sonraki sıcaklık tahmini."),
                        ("predicted_crowd_level", "LOW, MEDIUM, HIGH veya CRITICAL kalabalık tahmini."),
                        ("risk_score / risk_level", "0-100 risk skoru ve normal/warning/critical seviyesi."),
                        ("recommendation", "Operasyon panelinde gösterilecek karar destek önerisi."),
                        ("input_summary", "Tahmine giren IoT ve Video özet verileri."),
                    ],
                )
            ],
            endpoints=[
                "GET /health",
                "GET /model/status",
                "POST /model/train",
                "POST /predictions/generate",
                "GET /predictions/latest",
                "GET /predictions?zone={zone}&limit=50",
                "GET /zones/{zone}/risk",
            ],
            deployed_resources=[
                ("ECS Fargate service", "smart-city-final-ml-service API endpointlerini çalıştırır."),
                ("ECR repository", "smart-city-final/ml-service image versiyonlarını saklar."),
                ("DynamoDB", "ml_predictions tablosunu barındırır."),
                ("Lambda inference", "smart-city-final-ml-inference beş bölge için tahmin üretip DynamoDB'ye yazar."),
                ("CloudFront + ALB", "/api/ml path'ini ML target group'una yönlendirir."),
                ("SageMaker alternatifi", "Daha ölçekli senaryoda joblib/scikit-learn modelinin endpoint olarak sunulabileceği seçenek."),
            ],
            validation=[
                ("CloudFront ML API health", "/api/ml/health status=ok ve model_ready=true döndürdü."),
                ("Prediction latest", "/api/ml/predictions/latest beş bölge için prediction döndürdü."),
                ("Lambda inference", "Manuel invoke sonucu ok=true, count=5 olarak doğrulandı."),
                ("Frontend smoke", "ML Analytics ekranında risk kartları ve beş bölge tahmini görüntülendi."),
                ("ECS service", "desired=1, running=1, target health=healthy."),
            ],
            challenges=[
                ("Sentetik veriyle model kararlılığı", "Scikit-learn modeli gerçek kullanıldı; risk kararı rule-based motor ile açıklanabilir tutuldu."),
                ("Veri sahipliği", "ML servisi IoT/Video tablolarını okumadı; REST API ile veri aldı."),
                ("SageMaker maliyeti", "Demo için Lambda inference seçildi; raporda SageMaker ölçekli alternatif olarak bırakıldı."),
                ("Açıklanabilirlik", "Risk skoru sıcaklık, trafik, hava kalitesi ve kalabalık göstergeleriyle ilişkilendirildi."),
            ],
            conclusion=(
                "ML Analytics Service, Proje 3 kapsamındaki makine öğrenmesi modeli geliştirme, bulutta inference, veriden bilgi "
                "çıkarma ve tahmin üretme hedeflerini karşılar. Servis, IoT ve Video modüllerinin ham operasyon verisini karar "
                "destek çıktısına dönüştürür ve bu çıktıyı ayrı DynamoDB tablosunda saklar."
            ),
            sources=common_sources + ["Scikit-learn, Pandas, AWS Lambda ve Amazon SageMaker dokümantasyonları."],
        ),
    ]


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    built = []
    for spec in specs():
        built.append(build_docx(spec))
        built.append(build_pdf(spec))
    print("\n".join(str(p) for p in built))


if __name__ == "__main__":
    main()
